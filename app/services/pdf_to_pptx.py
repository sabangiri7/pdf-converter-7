"""Convert a PDF into a PowerPoint deck (one slide per page as an image).

Pure service: no Flask. Renders pages with PyMuPDF, builds .pptx with
python-pptx. No LibreOffice required.
"""
import io
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches

from ..errors import ToolError
from ..helpers import detect_kind

DEFAULT_DPI = 150
DPI_MIN, DPI_MAX = 72, 300
# 914400 EMUs per inch (Office Open XML)
_EMU_PER_INCH = 914400


def validate_dpi(dpi) -> int:
    try:
        value = int(dpi)
    except (TypeError, ValueError):
        raise ToolError("Resolution must be a whole number between "
                        f"{DPI_MIN} and {DPI_MAX} dpi.")
    if isinstance(dpi, float) and dpi != value:
        raise ToolError("Resolution must be a whole number between "
                        f"{DPI_MIN} and {DPI_MAX} dpi.")
    if not (DPI_MIN <= value <= DPI_MAX):
        raise ToolError(f"Resolution must be between {DPI_MIN} and "
                        f"{DPI_MAX} dpi.")
    return value


def clean_options(form) -> dict:
    return {"dpi": validate_dpi(form.get("dpi", DEFAULT_DPI))}


def run(inputs: list[Path], output_dir: Path, *,
        dpi: int = DEFAULT_DPI) -> Path:
    """Render each page of the single PDF onto its own PowerPoint slide."""
    dpi = validate_dpi(dpi)
    inputs = [Path(p) for p in (inputs or [])]
    if len(inputs) != 1:
        raise ToolError("This tool converts one PDF at a time. Please upload "
                        "a single PDF file.")
    src = inputs[0]
    if not src.is_file():
        raise ToolError("Your uploaded file is missing or expired. Please "
                        "upload it again.")
    if detect_kind(src) != "pdf":
        raise ToolError("That file doesn't look like a valid PDF. Please "
                        "upload a PDF file.")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / "converted.pptx"

    try:
        doc = pymupdf.open(src)
    except Exception:
        doc = None
    if doc is None:
        raise ToolError("That file doesn't look like a valid PDF, or it is "
                        "corrupted. Please try a different file.")
    try:
        try:
            if doc.needs_pass:
                raise ToolError("That PDF is password-protected. Please "
                                "remove the password and try again.")
            n_pages = doc.page_count
        except ToolError:
            raise
        except Exception:
            raise ToolError("That file doesn't look like a valid PDF, or it "
                            "is corrupted. Please try a different file.")
        if n_pages == 0:
            raise ToolError("That PDF has no pages to convert.")

        prs = Presentation()
        # Blank layout (no title/body placeholders).
        blank = prs.slide_layouts[6]
        slide_w = slide_h = None

        for i in range(n_pages):
            try:
                pix = doc[i].get_pixmap(dpi=dpi)
                png = pix.tobytes("png")
            except Exception:
                raise ToolError(
                    f"Page {i + 1} of the PDF is too large or damaged to "
                    f"render at {dpi} dpi. Please try a lower resolution.")

            width_in = pix.width / float(dpi)
            height_in = pix.height / float(dpi)
            if slide_w is None:
                # Size the deck to the first page; later pages are fitted.
                slide_w = Inches(width_in)
                slide_h = Inches(height_in)
                prs.slide_width = slide_w
                prs.slide_height = slide_h

            slide = prs.slides.add_slide(blank)
            left, top, pic_w, pic_h = _fit_box(
                width_in, height_in,
                prs.slide_width / _EMU_PER_INCH,
                prs.slide_height / _EMU_PER_INCH,
            )
            slide.shapes.add_picture(
                io.BytesIO(png),
                Inches(left), Inches(top),
                width=Inches(pic_w), height=Inches(pic_h),
            )

        try:
            prs.save(out)
        except Exception:
            out.unlink(missing_ok=True)
            raise ToolError("The PowerPoint file could not be written. "
                            "Please try again.")
        return out
    finally:
        doc.close()


def _fit_box(img_w: float, img_h: float,
             slide_w: float, slide_h: float) -> tuple[float, float, float, float]:
    """Return (left, top, width, height) in inches to fit image on slide."""
    scale = min(slide_w / img_w, slide_h / img_h) if img_w and img_h else 1.0
    pic_w = img_w * scale
    pic_h = img_h * scale
    left = (slide_w - pic_w) / 2.0
    top = (slide_h - pic_h) / 2.0
    return left, top, pic_w, pic_h
