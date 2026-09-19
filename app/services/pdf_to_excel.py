"""Extract tables (or text) from a PDF into an .xlsx workbook."""
from pathlib import Path

import pdfplumber
from openpyxl import Workbook

from app.errors import ToolError
from app.helpers import detect_kind

_BAD = ("That file doesn't look like a valid PDF, or it is corrupted. "
        "Please try a different file.")


def run(inputs: list[Path], output_dir: Path, **_options) -> Path:
    if len(inputs) != 1:
        raise ToolError("PDF to Excel works on one PDF at a time.")
    src = Path(inputs[0])
    if detect_kind(src) != "pdf":
        raise ToolError("Please upload a PDF file.")
    wb = Workbook()
    default = wb.active
    found = 0
    try:
        with pdfplumber.open(str(src)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables() or []
                if not tables:
                    text = (page.extract_text() or "").strip()
                    if not text:
                        continue
                    ws = wb.create_sheet(title=f"Page {i}"[:31])
                    for row_i, line in enumerate(text.splitlines(), 1):
                        ws.cell(row=row_i, column=1, value=line)
                    found += 1
                    continue
                for t_i, table in enumerate(tables, 1):
                    name = f"Page {i}" if len(tables) == 1 else f"P{i}-T{t_i}"
                    ws = wb.create_sheet(title=name[:31])
                    for r, row in enumerate(table or [], 1):
                        for c, cell in enumerate(row or [], 1):
                            ws.cell(row=r, column=c,
                                    value="" if cell is None else str(cell))
                    found += 1
    except ToolError:
        raise
    except Exception:
        raise ToolError(_BAD)
    if found == 0:
        raise ToolError("No tables or selectable text were found to export.")
    if default.title in wb.sheetnames and len(wb.sheetnames) > 1:
        wb.remove(default)
    out = Path(output_dir) / "tables.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except Exception:
        raise ToolError("The Excel file could not be written. Please try again.")
    return out
