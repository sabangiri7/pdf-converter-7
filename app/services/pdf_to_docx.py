"""Convert a text PDF into a Word .docx (one paragraph block per page).

Uses pdfplumber for extraction. Scanned/image-only PDFs produce little or
no text — callers should OCR first. Pure service: no Flask.
"""
from pathlib import Path

import pdfplumber
from docx import Document

from ..errors import ToolError
from ..helpers import detect_kind, page_count

_MAX_PAGES = 200


def run(inputs: list[Path], output_dir: Path, **_options) -> Path:
    """Extract text from one PDF into a .docx document."""
    inputs = [Path(p) for p in (inputs or [])]
    if len(inputs) != 1:
        raise ToolError("This tool converts one PDF at a time. Please upload "
                        "a single PDF file.")
    src = inputs[0]
    if not src.is_file():
        raise ToolError("Your uploaded file is missing or expired. Please "
                        "upload it again.")
    if detect_kind(src) != "pdf":
        raise ToolError("That file doesn't look like a valid PDF.")

    total = page_count(src)
    if total < 1:
        raise ToolError("That PDF does not contain any pages.")
    if total > _MAX_PAGES:
        raise ToolError(f"That PDF has too many pages (limit {_MAX_PAGES}).")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / "converted.docx"

    doc = Document()
    extracted_any = False
    try:
        with pdfplumber.open(src) as pdf:
            for i, page in enumerate(pdf.pages):
                if i > 0:
                    doc.add_page_break()
                text = (page.extract_text() or "").strip()
                if text:
                    extracted_any = True
                    for block in text.split("\n"):
                        doc.add_paragraph(block)
                else:
                    doc.add_paragraph(f"[Page {i + 1}: no selectable text]")
    except ToolError:
        raise
    except Exception:
        raise ToolError("That PDF could not be read. It may be corrupted or "
                        "encrypted.")

    if not extracted_any:
        raise ToolError(
            "No selectable text was found in that PDF. If it is a scan, "
            "run OCR PDF first, then try again.")

    try:
        doc.save(out)
    except Exception:
        out.unlink(missing_ok=True)
        raise ToolError("The Word document could not be written. Please try "
                        "again.")
    return out
