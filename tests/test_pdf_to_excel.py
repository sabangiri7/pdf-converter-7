"""Tests for pdf_to_excel."""
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.errors import ToolError
from app.services.pdf_to_excel import run


def _upload(pdf_path: Path, filename="x.pdf"):
    return {"files": [(io.BytesIO(pdf_path.read_bytes()), filename)]}


def test_service_ok(sample_pdf_2p, tmp_path):
    out = run([sample_pdf_2p], tmp_path)
    assert out.is_file() and out.suffix == ".xlsx"
    wb = load_workbook(out)
    assert len(wb.sheetnames) >= 1


def test_service_rejects_bad_input(tmp_path):
    junk = tmp_path / "junk.pdf"
    junk.write_bytes(b"not a pdf")
    with pytest.raises(ToolError):
        run([junk], tmp_path)


def test_get_page(client):
    r = client.get("/t/pdf_to_excel")
    assert r.status_code == 200
    assert b"PDF to Excel" in r.data


def test_route_roundtrip(client, sample_pdf_2p):
    r = client.post("/t/pdf_to_excel", data=_upload(sample_pdf_2p),
                    content_type="multipart/form-data")
    assert r.status_code == 200
    assert b"/dl/" in r.data
