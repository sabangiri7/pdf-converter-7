"""Tests for csv_to_xlsx: pure service + Flask route."""
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.errors import ToolError
from app.helpers import detect_kind, looks_like_csv
from app.services.csv_to_xlsx import run

ROOT = Path(__file__).resolve().parent.parent


def _csv(tmp_path, text: str, name="data.csv") -> Path:
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def test_looks_like_csv_helpers():
    assert looks_like_csv(b"a,b,c\n1,2,3\n")
    assert looks_like_csv(b"name;age\nAda;36\n")
    assert not looks_like_csv(b"hello world")
    assert not looks_like_csv(b"%PDF-1.4 fake")
    assert not looks_like_csv(b"\x00binary")


def test_detect_kind_csv(tmp_path):
    p = _csv(tmp_path, "x,y\n1,2\n")
    assert detect_kind(p) == "csv"


def test_happy_path(tmp_path):
    src = _csv(tmp_path, "name,score\nAda,10\nBob,9\n")
    out = run([src], tmp_path / "out")
    assert out.suffix == ".xlsx"
    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("name", "score")
    assert rows[1] == ("Ada", "10")
    assert rows[2] == ("Bob", "9")


def test_semicolon_delimiter(tmp_path):
    src = _csv(tmp_path, "a;b\n1;2\n")
    out = run([src], tmp_path / "out")
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert rows[0] == ("a", "b")


def test_rejects_non_csv(sample_image_png, tmp_path):
    with pytest.raises(ToolError):
        run([sample_image_png], tmp_path)


def test_requires_single(tmp_path):
    a = _csv(tmp_path, "a,b\n1,2\n", "a.csv")
    b = _csv(tmp_path, "c,d\n3,4\n", "b.csv")
    with pytest.raises(ToolError):
        run([a, b], tmp_path / "o")


def test_service_pure_no_flask():
    src = (ROOT / "app" / "services" / "csv_to_xlsx.py").read_text("utf-8")
    for token in ("from flask", "import flask", "Blueprint"):
        assert token not in src


def test_get_page(client):
    r = client.get("/t/csv_to_xlsx/")
    assert r.status_code == 200
    assert b"CSV" in r.data or b"csv" in r.data


def test_route_roundtrip(client, tmp_path):
    raw = b"col1,col2\nx,y\n"
    r = client.post(
        "/t/csv_to_xlsx/",
        data={"files": [(io.BytesIO(raw), "sheet.csv")]},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert r.status_code == 200
    assert b"/dl/" in r.data
    assert b".xlsx" in r.data


def test_route_rejects_non_csv(client, sample_image_jpg):
    r = client.post(
        "/t/csv_to_xlsx/",
        data={"files": [(open(sample_image_jpg, "rb"), "x.jpg")]},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert r.status_code == 400
